from io import BytesIO
from pathlib import Path
from threading import RLock
import gc

import numpy as np
import pandas as pd
import pymupdf as fitz
from PIL import Image
from docx import Document

_OCR = None
_OCR_LOCK = RLock()


def _ocr_engine():
    global _OCR
    if _OCR is not None:
        return _OCR
    with _OCR_LOCK:
        if _OCR is None:
            from rapidocr import RapidOCR
            _OCR = RapidOCR(
                params={
                    "Global.log_level": "warning",
                    "Global.use_cls": False,
                    "Global.max_side_len": 736,
                    "EngineConfig.onnxruntime.intra_op_num_threads": 1,
                    "EngineConfig.onnxruntime.inter_op_num_threads": 1,
                    "EngineConfig.onnxruntime.enable_cpu_mem_arena": False,
                    "Det.limit_side_len": 736,
                    "Det.limit_type": "max",
                    "Rec.rec_batch_num": 1,
                }
            )
    return _OCR


def _ocr_image(image_bytes: bytes) -> str:
    try:
        with Image.open(BytesIO(image_bytes)) as source:
            image = source.convert("RGB")
            result = _ocr_engine()(
                np.asarray(image),
                use_cls=False,
            )

        if hasattr(result, "txts"):
            values = result.txts or ()
        elif isinstance(result, tuple) and result:
            values = getattr(result[0], "txts", None) or result[0]
        else:
            values = result or []

        if isinstance(values, str):
            return values.strip()
        return "\n".join(str(value) for value in values if value).strip()
    except Exception:
        return ""


def extract(path: Path, progress_callback=None):
    ext = path.suffix.lower()
    text = ""
    refs = []
    ocr = False
    pages = 0

    if ext == ".pdf":
        with fitz.open(path) as doc:
            pages = len(doc)
            native_pages: list[str] = []
            native_chars = 0

            # First pass: extract native PDF text without initializing OCR.
            # Many PDFs contain mostly vector/text content with an occasional
            # image-only page; OCR should not be loaded just because one page
            # is blank to PyMuPDF.
            for i, page in enumerate(doc):
                page_text = (page.get_text("text") or "").strip()
                native_pages.append(page_text)
                native_chars += len(page_text)

            native_coverage = (
                sum(1 for value in native_pages if value) / max(1, pages)
            )

            use_ocr_fallback = (
                pages > 0
                and native_chars < max(1000, pages * 80)
                and native_coverage < 0.5
            )

            for i, page in enumerate(doc):
                page_text = native_pages[i]

                if not page_text and use_ocr_fallback:
                    try:
                        # OCR each page at a moderate raster size to reduce
                        # transient memory pressure on small Render instances.
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(1.0, 1.0),
                            alpha=False,
                        )
                        image_bytes = pix.tobytes("png")
                        del pix
                        page_text = _ocr_image(image_bytes)
                        del image_bytes
                        gc.collect()
                        if page_text:
                            ocr = True
                    except Exception:
                        page_text = ""
                    finally:
                        gc.collect()

                if progress_callback:
                    try:
                        # Reserve 25-40% of the pipeline for document extraction
                        # so scanned PDFs show live page-level progress instead
                        # of appearing frozen at 25%.
                        progress_callback(
                            25 + int(15 * (i + 1) / max(1, pages))
                        )
                    except Exception:
                        pass

                if page_text:
                    text += (
                        f"\n\n[[SOURCE:page {i + 1}]]\n"
                        f"{page_text}\n"
                    )
                    refs.append(f"page {i + 1}")

    elif ext == ".docx":
        document = Document(path)
        paragraphs = [
            p.text.strip()
            for p in document.paragraphs
            if p.text.strip()
        ]
        tables = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        text = "\n".join(paragraphs + [row for row in tables if row])
        refs = ["document"]

    elif ext in {".xlsx", ".xls"}:
        sheets = pd.read_excel(path, sheet_name=None)
        frames = []
        for sheet_name, frame in sheets.items():
            if not frame.empty:
                frames.append(
                    f"\n\n[[SOURCE:sheet {sheet_name}]]\n"
                    + frame.to_csv(index=True)
                )
        text = "\n".join(frames)
        refs = [f"sheet {name}" for name in sheets] or ["table"]

    elif ext == ".csv":
        frame = pd.read_csv(path)
        text = "\n\n[[SOURCE:table]]\n" + frame.to_csv(index=True)
        refs = ["table"]

    elif ext in {".jpg", ".jpeg", ".png"}:
        text = _ocr_image(path.read_bytes())
        ocr = bool(text)
        refs = ["image"]

    elif ext in {".txt", ".md"}:
        text = path.read_text(errors="ignore")
        refs = ["text"]

    return text, refs, ocr, pages
